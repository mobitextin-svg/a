"""
ZyvoRCS Sender - Flask Backend v2
Zyvo Tools | app.py

New in v2:
  - /sims  endpoint — detect SIM slots
  - /paste_contacts — accept pasted numbers (no Excel needed)
  - /send  now accepts sim_mode, sim1_sub_id, sim2_sub_id
"""

import os
import json
import re
import threading
import subprocess
from flask import Flask, request, jsonify, render_template, send_file
import openpyxl

app = Flask(__name__)

# ─── Shared State (thread-safe) ──────────────────────────────────
status = {
    'running': False, 'paused': False, 'stop_flag': False,
    'sent': 0, 'failed': 0, 'skipped': 0, 'total': 0,
    'current': '', 'log': [], 'error': '', 'balance': None,
}
lock        = threading.Lock()
send_thread = None

# Cached paste contacts (numbers only mode)
paste_contacts_cache = []

# ─── Paths ───────────────────────────────────────────────────────
BASE_DIR      = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR    = os.path.join(BASE_DIR, 'uploads')
REPORT_DIR    = os.path.join(BASE_DIR, 'reports')
DATA_DIR      = os.path.join(BASE_DIR, 'data')
TEMPLATE_FILE = os.path.join(DATA_DIR, 'templates.json')

for d in [UPLOAD_DIR, REPORT_DIR, DATA_DIR]:
    os.makedirs(d, exist_ok=True)


# ═══════════════════════════════════════════════════════════════════
# ROUTES
# ═══════════════════════════════════════════════════════════════════

@app.route('/')
def index():
    return render_template('index.html')


# ─── Devices ─────────────────────────────────────────────────────

@app.route('/devices')
def get_devices():
    try:
        r = subprocess.run(['adb', 'devices'], capture_output=True, text=True, timeout=5)
        lines = r.stdout.strip().split('\n')[1:]
        devices = []
        for line in lines:
            if line.strip() and '\t' in line:
                serial, state = line.split('\t', 1)
                devices.append({'serial': serial.strip(), 'state': state.strip()})
        return jsonify({'devices': devices})
    except FileNotFoundError:
        return jsonify({'devices': [], 'error': 'adb not found. Add adb to PATH.'})
    except Exception as e:
        return jsonify({'devices': [], 'error': str(e)})


@app.route('/device_info', methods=['POST'])
def device_info():
    """Return phone model, brand, and Android version for a connected device"""
    device = (request.json or {}).get('device', '')
    from adb_engine import get_device_info
    info = get_device_info(device)
    return jsonify(info)


@app.route('/sims', methods=['POST'])
def get_sims():
    """Detect SIM slots on a connected device"""
    device = (request.json or {}).get('device', '')
    from adb_engine import get_sims
    sims = get_sims(device)
    return jsonify({'sims': sims})


@app.route('/pair', methods=['POST'])
def pair_device():
    data     = request.json
    ip       = data.get('ip', '').strip()
    port     = data.get('port', '37000').strip()
    password = data.get('password', '').strip()

    if not ip or not password:
        return jsonify({'success': False, 'message': 'IP and password required'})
    try:
        pr = subprocess.run(['adb', 'pair', f'{ip}:{port}', password],
                            capture_output=True, text=True, timeout=20)
        out = pr.stdout + pr.stderr
        if 'Successfully paired' not in out and 'successfully' not in out.lower():
            return jsonify({'success': False, 'message': out.strip() or 'Pair failed'})
        cr = subprocess.run(['adb', 'connect', f'{ip}:5555'],
                            capture_output=True, text=True, timeout=10)
        return jsonify({'success': True, 'message': cr.stdout.strip()})
    except subprocess.TimeoutExpired:
        return jsonify({'success': False, 'message': 'Timeout — check IP and try again'})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)})


@app.route('/connect', methods=['POST'])
def connect_device():
    ip = (request.json or {}).get('ip', '').strip()
    try:
        r = subprocess.run(['adb', 'connect', f'{ip}:5555'],
                           capture_output=True, text=True, timeout=10)
        return jsonify({'success': 'connected' in r.stdout.lower(),
                        'message': r.stdout.strip()})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)})


# ─── Contacts ────────────────────────────────────────────────────

@app.route('/upload', methods=['POST'])
def upload_contacts():
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': 'No file provided'})
    file = request.files['file']
    if not file.filename.endswith('.xlsx'):
        return jsonify({'success': False, 'message': 'Only .xlsx supported'})
    path = os.path.join(UPLOAD_DIR, 'contacts.xlsx')
    file.save(path)
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb.active
        headers   = [str(c.value) if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
        row_count = ws.max_row - 1
        wb.close()
        return jsonify({'success': True, 'columns': headers, 'rows': row_count})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)})


@app.route('/paste_contacts', methods=['POST'])
def paste_contacts():
    """
    Accept raw pasted numbers (one per line, or comma-separated).
    Returns count of valid numbers found.
    """
    global paste_contacts_cache
    data = request.json or {}
    raw  = data.get('text', '')

    # Extract all 10-digit (or 10-12 digit with country code) mobile numbers
    numbers = re.findall(r'(?:\+?91)?([6-9]\d{9})', raw.replace(' ', '').replace('-', ''))
    # Deduplicate while preserving order
    seen = set()
    unique = []
    for n in numbers:
        if n not in seen:
            seen.add(n)
            unique.append(n)

    paste_contacts_cache = [{'mobile': n, 'name': n} for n in unique]

    return jsonify({
        'success': True,
        'count':   len(unique),
        'preview': unique[:5],     # first 5 for UI confirmation
        'total':   len(unique),
    })


# ─── Templates ───────────────────────────────────────────────────

@app.route('/template', methods=['GET', 'POST'])
def template_handler():
    if request.method == 'POST':
        data = request.json
        t    = _load_templates()
        t[data['name']] = data['content']
        _save_templates(t)
        return jsonify({'success': True})
    return jsonify({'templates': _load_templates()})


@app.route('/template/<name>', methods=['DELETE'])
def delete_template(name):
    t = _load_templates()
    t.pop(name, None)
    _save_templates(t)
    return jsonify({'success': True})


@app.route('/templates/export')
def export_templates():
    """Download all templates as JSON file"""
    path = TEMPLATE_FILE
    if not os.path.exists(path):
        _save_templates({})
    return send_file(path, as_attachment=True, download_name='zyvorcs_templates.json')


@app.route('/templates/import', methods=['POST'])
def import_templates():
    """Import templates from uploaded JSON — merges with existing"""
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': 'No file provided'})
    file = request.files['file']
    try:
        imported = json.load(file)
        if not isinstance(imported, dict):
            return jsonify({'success': False, 'message': 'Invalid format — expected JSON object'})
        existing = _load_templates()
        existing.update(imported)   # merge, imported wins on conflict
        _save_templates(existing)
        return jsonify({'success': True, 'count': len(imported), 'total': len(existing)})
    except Exception as e:
        return jsonify({'success': False, 'message': f'Parse error: {e}'})


# ─── Send ─────────────────────────────────────────────────────────

@app.route('/send', methods=['POST'])
def start_send():
    global send_thread, paste_contacts_cache
    with lock:
        if status['running']:
            return jsonify({'success': False, 'message': 'Already running'})

    data        = request.json
    col_map     = data.get('column_map', {})
    template    = data.get('template', '')
    delay       = int(data.get('delay', 5))
    device      = data.get('device', '')
    sim_mode    = data.get('sim_mode', 'rotate')
    sim1_sub_id = int(data.get('sim1_sub_id', 1))
    sim2_sub_id = int(data.get('sim2_sub_id', 2))
    sim1_number = str(data.get('sim1_number', '')).strip()
    sim2_number = str(data.get('sim2_number', '')).strip()
    device_pin  = data.get('device_pin', '')
    daily_limit = int(data.get('daily_limit', 100))
    min_delay   = int(data.get('min_delay', 60))
    max_delay   = int(data.get('max_delay', 180))
    if min_delay > max_delay:          # swap if user entered them backwards
        min_delay, max_delay = max_delay, min_delay
    batch_size  = max(0, int(data.get('batch_size', 20)))
    batch_pause = max(5, int(data.get('batch_pause', 300)))
    hours_start   = min(23, max(0, int(data.get('hours_start', 9))))
    hours_end     = min(24, max(0, int(data.get('hours_end', 20))))
    hours_enforce = bool(int(data.get('hours_enforce', 1)))
    source      = data.get('source', 'excel')    # 'excel' | 'paste'

    if not template:
        return jsonify({'success': False, 'message': 'Message template is empty'})

    # Load contacts from correct source
    if source == 'paste':
        contacts = list(paste_contacts_cache)
    else:
        contacts = _load_contacts(col_map)

    if not contacts:
        return jsonify({'success': False,
                        'message': 'No contacts. Upload Excel or paste numbers first.'})

    with lock:
        status.update({
            'running': True, 'paused': False, 'stop_flag': False,
            'sent': 0, 'failed': 0, 'skipped': 0, 'total': len(contacts),
            'current': '', 'log': [], 'error': '', 'balance': None,
        })

    from adb_engine import run_bulk_send
    send_thread = threading.Thread(
        target=run_bulk_send,
        args=(contacts, template, delay, device, status, lock, REPORT_DIR),
        kwargs={
            'sim_mode':    sim_mode,
            'sim1_sub_id': sim1_sub_id,
            'sim2_sub_id': sim2_sub_id,
            'device_pin':  device_pin,
            'daily_limit': daily_limit,
            'min_delay':   min_delay,
            'max_delay':   max_delay,
            'batch_size':  batch_size,
            'batch_pause': batch_pause,
            'sim1_number': sim1_number,
            'sim2_number': sim2_number,
            'hours_start':   hours_start,
            'hours_end':     hours_end,
            'hours_enforce': hours_enforce,
        },
        daemon=True
    )
    send_thread.start()
    return jsonify({'success': True, 'total': len(contacts)})


@app.route('/status')
def get_status():
    with lock:
        return jsonify({
            'running': status['running'], 'paused': status['paused'],
            'sent':    status['sent'],    'failed': status['failed'],
            'skipped': status.get('skipped', 0),
            'total':   status['total'],   'current': status['current'],
            'log':     status['log'][-80:],
            'error':   status['error'],
            'balance': status.get('balance'),
        })


@app.route('/pause', methods=['POST'])
def pause_send():
    with lock:
        status['paused'] = not status['paused']
        paused = status['paused']
    return jsonify({'paused': paused})


@app.route('/stop', methods=['POST'])
def stop_send():
    with lock:
        status['stop_flag'] = True
    return jsonify({'success': True})


@app.route('/test_send', methods=['POST'])
def test_send():
    """
    Diagnostic: runs one test send step-by-step and returns
    per-step results so you can see exactly where it fails.
    """
    data    = request.json or {}
    device  = data.get('device', '')
    mobile  = data.get('mobile', '9876543210')
    message = data.get('message', 'ZyvoRCS test message')
    device_pin = data.get('device_pin', '')

    from adb_engine import run_diagnostic
    steps = run_diagnostic(device, mobile, message, device_pin=device_pin)
    return jsonify({'steps': steps})


@app.route('/dump_ui', methods=['POST'])
def dump_ui():
    """Dump current UI hierarchy from device — for debugging send button location"""
    device = (request.json or {}).get('device', '')
    from adb_engine import get_ui_dump
    xml = get_ui_dump(device)
    return jsonify({'xml': xml})


@app.route('/report')
def download_report():
    path = os.path.join(REPORT_DIR, 'report.xlsx')
    if os.path.exists(path):
        return send_file(path, as_attachment=True, download_name='ZyvoRCS_Report.xlsx')
    return jsonify({'error': 'No report yet'}), 404


# ═══════════════════════════════════════════════════════════════════
# HELPERS
# ═══════════════════════════════════════════════════════════════════

def _load_templates():
    if os.path.exists(TEMPLATE_FILE):
        with open(TEMPLATE_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {}

def _save_templates(t):
    with open(TEMPLATE_FILE, 'w', encoding='utf-8') as f:
        json.dump(t, f, ensure_ascii=False, indent=2)

def _load_contacts(col_map):
    path = os.path.join(UPLOAD_DIR, 'contacts.xlsx')
    if not os.path.exists(path):
        return []
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb.active
        headers  = [str(c.value) if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
        contacts = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            contact = {}
            for var, col_name in col_map.items():
                if col_name in headers:
                    idx = headers.index(col_name)
                    contact[var] = str(row[idx]) if row[idx] is not None else ''
            if contact.get('mobile', '').strip():
                contacts.append(contact)
        wb.close()
        return contacts
    except Exception:
        return []


if __name__ == '__main__':
    app.run(host='127.0.0.1', port=5000, debug=False, use_reloader=False)
