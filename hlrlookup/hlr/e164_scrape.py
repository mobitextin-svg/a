"""
Shared core for looking up numbers on e164.com with a real browser (Playwright).

Both the CLI (playwright_lookup.py) and the web app (e164_app.py) use this so
there is exactly one place that knows how to drive e164.com and parse its
result block. Everything here returns the REAL fields e164.com shows — there is
no synthetic/mock data in this module.

e164.com result fields (in display order):
  Prefix, Calling Code, ISO3, TADIG, MCCMNC, Type, Location,
  Operator Brand, Operator Company, Operator Group,
  Total Length Min, Total Length Max, Weight, Source
"""

import os
import re
import csv
import io

try:
    from .india import normalize_msisdn, to_e164
except Exception:  # standalone import fallback
    def normalize_msisdn(raw):
        d = re.sub(r"\D", "", raw or "")
        if d.startswith("0091"): d = d[4:]
        elif d.startswith("91") and len(d) == 12: d = d[2:]
        elif d.startswith("0") and len(d) == 11: d = d[1:]
        return d if len(d) == 10 and d[0] in "6789" else None
    def to_e164(n): return "+91" + n


E164_FIELDS = [
    "Prefix", "Calling Code", "ISO3", "TADIG", "MCCMNC", "Type", "Location",
    "Operator Brand", "Operator Company", "Operator Group",
    "Total Length Min", "Total Length Max", "Weight", "Source",
]

# Columns used for CSV / Excel output (input first so you can trace the row).
COLUMNS = ["input", "msisdn", "typed"] + E164_FIELDS + ["status", "error", "checked_at"]

SEARCH_GUESSES = [
    "input[type=search]", "input[type=tel]",
    "input[name*=number i]", "input[id*=number i]",
    "input[name*=phone i]", "input[id*=phone i]",
    "input[name*=search i]", "input[id*=search i]",
    "input[placeholder*=number i]", "input[placeholder*=phone i]",
    "input[type=text]:visible", "input:visible",
]


# ── parsing ─────────────────────────────────────────────────────────────────
def parse_e164_result(body_text):
    """Extract the labelled e164.com fields from the page's visible text."""
    fields = {}
    for label in E164_FIELDS:
        m = re.search(rf"(?m)^\s*{re.escape(label)}\s*:\s*(.+?)\s*$", body_text)
        if m:
            fields[label] = m.group(1).strip()
    return fields


def shown_number(body_text):
    """Digits from the 'Results for: …' line, or '' if absent."""
    m = re.search(r"Results for:\s*([0-9+\s]+)", body_text)
    return re.sub(r"\D", "", m.group(1)) if m else ""


def format_typed(raw, number_format="cc"):
    """
    Decide what to type into the search box.
    Returns (typed, national10, invalid_reason). invalid_reason is None if OK.
    """
    if number_format == "raw":
        return re.sub(r"\s+", "", raw), re.sub(r"\D", "", raw)[-10:], None
    n = normalize_msisdn(raw)
    if n is None:
        return None, None, "invalid Indian mobile number"
    if number_format == "e164":
        return to_e164(n), n, None
    if number_format == "national":
        return n, n, None
    return "91" + n, n, None   # cc: 91XXXXXXXXXX (matches e164.com)


# ── Excel input ─────────────────────────────────────────────────────────────
def read_numbers(path, column=""):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    col_idx = None
    if column:
        col = column.strip()
        if col.lower() in header:
            col_idx = header.index(col.lower())
        elif re.fullmatch(r"[A-Za-z]", col):
            col_idx = ord(col.upper()) - ord("A")
        elif col.isdigit():
            col_idx = int(col)
    if col_idx is None:
        for i, h in enumerate(header):
            if any(k in h for k in ("number", "phone", "msisdn", "mobile", "cell")):
                col_idx = i
                break
        if col_idx is None:
            col_idx = 0
    header_is_label = not re.search(r"\d", str(rows[0][col_idx] or ""))
    body = rows[1:] if header_is_label else rows
    out = []
    for r in body:
        if col_idx < len(r) and r[col_idx] not in (None, ""):
            out.append(str(r[col_idx]).strip())
    return out


def split_text(blob):
    """Numbers from pasted text (newline and/or comma separated)."""
    out = []
    for line in (blob or "").replace("\r", "\n").split("\n"):
        for part in line.split(","):
            part = part.strip().strip('"').strip("'")
            if part:
                out.append(part)
    return out


# ── the browser driver ──────────────────────────────────────────────────────
class E164Scraper:
    """
    Drives a single real browser page against e164.com. Use as a context
    manager and call ``lookup()`` per number. Sequential by design — one browser,
    polite to the site.
    """

    def __init__(self, url="https://www.e164.com", search_selector="",
                 submit_selector="", headless=True, timeout=30000, settle=600,
                 reload_each=False, executable_path="", slowmo=0):
        self.url = url
        self.search_selector = search_selector
        self.submit_selector = submit_selector
        self.headless = headless
        self.timeout = timeout
        self.settle = settle
        self.reload_each = reload_each
        self.executable_path = executable_path or os.environ.get("PLAYWRIGHT_CHROMIUM_EXECUTABLE", "")
        self.slowmo = slowmo
        self._pw = self._browser = self._ctx = self.page = None

    def __enter__(self):
        self.start()
        return self

    def __exit__(self, *exc):
        self.close()

    def start(self):
        from playwright.sync_api import sync_playwright
        self._pw = sync_playwright().start()
        kwargs = dict(headless=self.headless, slow_mo=self.slowmo)
        if self.executable_path:
            kwargs["executable_path"] = self.executable_path
        self._browser = self._pw.chromium.launch(**kwargs)
        self._ctx = self._browser.new_context()
        self.page = self._ctx.new_page()
        self.page.goto(self.url, wait_until="domcontentloaded", timeout=self.timeout)

    def _find_box(self):
        page = self.page
        if self.search_selector:
            return page.locator(self.search_selector).first
        for guess in SEARCH_GUESSES:
            loc = page.locator(guess).first
            try:
                if loc.count() and loc.is_visible():
                    return loc
            except Exception:
                continue
        return None

    def lookup(self, typed, national):
        """Search one number. Returns (fields_dict, error_str)."""
        import time
        from playwright.sync_api import TimeoutError as PWTimeout  # noqa: F401
        page = self.page
        if self.reload_each:
            page.goto(self.url, wait_until="domcontentloaded", timeout=self.timeout)

        box = self._find_box()
        if box is None:
            return {}, "search box not found (set search_selector)"

        try:
            box.click()
            box.fill("")
            box.press_sequentially(typed, delay=15)
        except Exception as exc:
            return {}, f"could not type into search box: {exc}"

        if self.submit_selector:
            try:
                page.locator(self.submit_selector).first.click(timeout=3000)
            except Exception:
                pass
        else:
            try:
                box.press("Enter")
            except Exception:
                pass

        deadline = time.time() + self.timeout / 1000.0
        body = ""
        while time.time() < deadline:
            body = page.inner_text("body")
            if national and national in shown_number(body):
                break
            page.wait_for_timeout(200)
        else:
            return {}, "timed out waiting for this number's result"

        page.wait_for_timeout(self.settle)
        body = page.inner_text("body")
        fields = parse_e164_result(body)
        if not fields:
            return {}, "results block not found / no fields parsed"
        return fields, ""

    def close(self):
        try:
            if self._browser:
                self._browser.close()
        finally:
            if self._pw:
                self._pw.stop()
            self._pw = self._browser = self._ctx = self.page = None


# ── output helpers ──────────────────────────────────────────────────────────
def rows_to_csv(rows):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=COLUMNS, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        w.writerow(r)
    return buf.getvalue()


def rows_to_txt(rows):
    out = ["Number lookup results — source: e164.com", ""]
    for r in rows:
        out.append("=" * 56)
        typed = r.get("typed") or ""
        out.append(f"Input: {r.get('input','')}" + (f"   (typed: {typed})" if typed else ""))
        if r.get("status") == "INVALID":
            out.append("  INVALID — not a valid Indian mobile number")
        elif r.get("error"):
            out.append(f"  ERROR — {r['error']}")
        else:
            for label in E164_FIELDS:
                if r.get(label):
                    out.append(f"  {label}: {r[label]}")
        out.append(f"  Checked: {r.get('checked_at','')}")
    return "\n".join(out) + "\n"
